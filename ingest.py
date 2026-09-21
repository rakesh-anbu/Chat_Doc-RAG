import os
import json
import uuid
import time
import zipfile
import logging
from typing import List, Dict, Any, Callable, Optional
from pathlib import Path
from dotenv import load_dotenv

# Parsers
from pypdf import PdfReader
import docx
from pptx import Presentation
from bs4 import BeautifulSoup
from striprtf.striprtf import rtf_to_text
import openpyxl
import yaml
import ebooklib
from ebooklib import epub

from google import genai
from google.genai import types
from pinecone import Pinecone

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

load_dotenv()

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
PINECONE_API_KEY = os.getenv("PINECONE_API_KEY")
PINECONE_INDEX_HOST = os.getenv("PINECONE_INDEX_HOST")

if not GEMINI_API_KEY or not PINECONE_API_KEY or not PINECONE_INDEX_HOST:
    raise ValueError("Missing environment variables. Check your .env file.")

genai_client = genai.Client(api_key=GEMINI_API_KEY)
pc = Pinecone(api_key=PINECONE_API_KEY)
index = pc.Index(host=PINECONE_INDEX_HOST)

EMBEDDING_MODEL = "gemini-embedding-001"

SUPPORTED_EXTENSIONS = {
    ".pdf", ".json", ".txt", ".md", ".csv",
    ".docx", ".doc", ".dot", ".dotx",
    ".pptx", ".ppt",
    ".xlsx", ".xls",
    ".yaml", ".yml", ".xml",
    ".epub",
    ".rtf", ".html", ".htm",
    ".hwpx", ".hwp"
}


def extract_text_from_pdf(pdf_path: str) -> str:
    reader = PdfReader(pdf_path)
    text_parts = []
    for page in reader.pages:
        try:
            extracted = page.extract_text()
            if extracted:
                text_parts.append(extracted.strip())
        except Exception:
            continue
    return "\n\n".join(text_parts)


def extract_text_from_json(json_path: str) -> str:
    with open(json_path, "r", encoding="utf-8", errors="ignore") as f:
        data = json.load(f)
    if isinstance(data, list):
        items = [json.dumps(item, ensure_ascii=False) if isinstance(item, (dict, list)) else str(item) for item in data]
        return "\n\n".join(items)
    elif isinstance(data, dict):
        return json.dumps(data, indent=2, ensure_ascii=False)
    else:
        return str(data)


def extract_text_from_docx(doc_path: str) -> str:
    doc = docx.Document(doc_path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)
    return "\n\n".join(paragraphs)


def extract_text_from_pptx(ppt_path: str) -> str:
    prs = Presentation(ppt_path)
    text_runs = []
    for slide_idx, slide in enumerate(prs.slides, start=1):
        slide_texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_texts.append(shape.text.strip())
        if slide_texts:
            text_runs.append(f"[Slide {slide_idx}]\n" + "\n".join(slide_texts))
    return "\n\n".join(text_runs)


def extract_text_from_excel(excel_path: str) -> str:
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheets_text = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows_text = []
        for row in ws.iter_rows(values_only=True):
            row_vals = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
            if row_vals:
                rows_text.append(" | ".join(row_vals))
        if rows_text:
            sheets_text.append(f"[Sheet: {sheet}]\n" + "\n".join(rows_text))
    return "\n\n".join(sheets_text)


def extract_text_from_epub(epub_path: str) -> str:
    book = epub.read_epub(epub_path)
    chapters = []
    for item in book.get_items():
        if item.get_type() == ebooklib.ITEM_DOCUMENT:
            soup = BeautifulSoup(item.get_content(), "html.parser")
            text = soup.get_text(separator="\n\n").strip()
            if text:
                chapters.append(text)
    return "\n\n".join(chapters)


def extract_text_from_yaml(yaml_path: str) -> str:
    with open(yaml_path, "r", encoding="utf-8", errors="ignore") as f:
        data = yaml.safe_load(f)
        return json.dumps(data, indent=2, ensure_ascii=False) if isinstance(data, (dict, list)) else str(data)


def extract_text_from_html(html_path: str) -> str:
    with open(html_path, "r", encoding="utf-8", errors="ignore") as f:
        soup = BeautifulSoup(f.read(), "html.parser")
        for script in soup(["script", "style"]):
            script.decompose()
        return soup.get_text(separator="\n\n").strip()


def extract_text_from_rtf(rtf_path: str) -> str:
    with open(rtf_path, "r", encoding="utf-8", errors="ignore") as f:
        return rtf_to_text(f.read())


def extract_text_from_hwpx(hwpx_path: str) -> str:
    text_list = []
    try:
        with zipfile.ZipFile(hwpx_path, 'r') as z:
            for filename in z.namelist():
                if filename.startswith('Contents/section') and filename.endswith('.xml'):
                    with z.open(filename) as xml_file:
                        soup = BeautifulSoup(xml_file.read(), 'xml')
                        text_list.append(soup.get_text(separator=' '))
    except Exception as e:
        logger.warning(f"Error reading hwpx package: {e}")
    return "\n\n".join(text_list)


def extract_text_universal(file_path: str) -> str:
    path = Path(file_path)
    ext = path.suffix.lower()
    
    if ext == ".pdf":
        return extract_text_from_pdf(str(path))
    elif ext == ".json":
        return extract_text_from_json(str(path))
    elif ext in {".docx", ".dotx", ".doc", ".dot"}:
        try:
            return extract_text_from_docx(str(path))
        except Exception:
            with open(str(path), "rb") as f:
                content = f.read()
                return "".join(chr(b) for b in content if 32 <= b <= 126 or b in (10, 13))
    elif ext in {".pptx", ".ppt"}:
        return extract_text_from_pptx(str(path))
    elif ext in {".xlsx", ".xls"}:
        return extract_text_from_excel(str(path))
    elif ext == ".epub":
        return extract_text_from_epub(str(path))
    elif ext in {".yaml", ".yml"}:
        return extract_text_from_yaml(str(path))
    elif ext in {".html", ".htm", ".xml"}:
        return extract_text_from_html(str(path))
    elif ext == ".rtf":
        return extract_text_from_rtf(str(path))
    elif ext in {".hwpx", ".hwp"}:
        try:
            return extract_text_from_hwpx(str(path))
        except Exception:
            with open(str(path), "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
    elif ext in {".txt", ".md", ".csv"}:
        with open(str(path), "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    else:
        with open(str(path), "r", encoding="utf-8", errors="ignore") as f:
            return f.read()


def chunk_text(text: str, chunk_size: int = 800, chunk_overlap: int = 60) -> List[str]:
    words = text.split()
    if not words:
        return []
    
    chunks = []
    start = 0
    step = max(1, chunk_size - chunk_overlap)
    total_words = len(words)
    
    while start < total_words:
        end = start + chunk_size
        chunk = " ".join(words[start:end]).strip()
        if len(chunk) > 10:
            chunks.append(chunk)
        if end >= total_words:
            break
        start += step
    return chunks


def get_embeddings_fast(texts: List[str], max_retries: int = 4) -> List[List[float]]:
    if not texts:
        return []
        
    for attempt in range(max_retries):
        try:
            response = genai_client.models.embed_content(
                model=EMBEDDING_MODEL,
                contents=texts,
                config=types.EmbedContentConfig(
                    task_type="RETRIEVAL_DOCUMENT",
                    output_dimensionality=768
                )
            )
            return [e.values for e in response.embeddings]
        except Exception as e:
            err_str = str(e)
            wait_time = 1.5 * (attempt + 1)
            logger.warning(f"Embedding attempt {attempt + 1}/{max_retries} failed ({err_str[:60]}). Retrying in {wait_time}s...")
            time.sleep(wait_time)
            
    raise RuntimeError(f"Critical: Failed to generate embeddings after {max_retries} attempts.")


def ingest_file(file_path: str, batch_size: int = 96, progress_callback: Optional[Callable[[float], None]] = None, *args, **kwargs):
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    logger.info(f"Starting ultra-fast ingestion for: {path.name}")
    raw_text = extract_text_universal(str(path))
    
    if not raw_text or not raw_text.strip():
        logger.warning(f"No readable text content extracted from {path.name}")
        return

    chunks = chunk_text(raw_text, chunk_size=800, chunk_overlap=60)
    total_chunks = len(chunks)
    logger.info(f"Generated {total_chunks} text chunks for {path.name}.")

    if total_chunks == 0:
        return

    for i in range(0, total_chunks, batch_size):
        batch_chunks = chunks[i:i + batch_size]
        embeddings = get_embeddings_fast(batch_chunks)
        
        vectors = []
        for idx, (chunk, emb) in enumerate(zip(batch_chunks, embeddings)):
            vector_id = f"{path.stem}_chunk_{i + idx}_{uuid.uuid4().hex[:6]}"
            vectors.append({
                "id": vector_id,
                "values": emb,
                "metadata": {
                    "source": path.name,
                    "chunk_index": i + idx,
                    "text": chunk
                }
            })
        
        index.upsert(vectors=vectors)
            
        if progress_callback:
            try:
                progress_callback(min(1.0, (i + len(batch_chunks)) / total_chunks))
            except Exception:
                pass

    logger.info(f"SUCCESS: Ingestion of {path.name} ({total_chunks} chunks) completed!")


if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        print("Usage: python ingest.py <path_to_document>")
        sys.exit(1)
    
    ingest_file(sys.argv[1])
